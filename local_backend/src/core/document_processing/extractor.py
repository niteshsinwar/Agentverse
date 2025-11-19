"""
Document Extractor Module
Handles extraction of text content from various file formats
Returns (content, modality, metadata) tuple for RAG pipeline
"""
from __future__ import annotations

from pathlib import Path
from typing import Tuple, Dict, Any, Optional, List
import os
import tempfile

import pandas as pd
from PIL import Image
import fitz  # PyMuPDF for PDF and image processing
import pypdf
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

MAX_TABULAR_SAMPLE_ROWS = 2000
TABULAR_PREVIEW_ROWS = 10
NUMERIC_SUMMARY_MAX_COLS = 10
MAX_COLUMN_NAMES = 40


class DocumentExtractor:
    """Extracts text content from various document formats"""

    def __init__(self, settings=None):
        from src.core.config.settings import get_settings
        self.settings = settings or get_settings()
        self._settings_signature = id(self.settings)

        # Build supported formats from global settings
        self._format_mapping = {
            'pdf': self._extract_pdf_content,
            'docx': self._extract_docx_content,
            'pptx': self._extract_pptx_content,
            'csv': self._extract_csv_content,
            'txt': self._extract_text_content,
            'md': self._extract_text_content,
            'json': self._extract_text_content,
            'xml': self._extract_text_content,
            'html': self._extract_text_content,
            'htm': self._extract_text_content,
            'xlsx': self._extract_spreadsheet_content,
            'xls': self._extract_spreadsheet_content,
            'rtf': self._extract_text_content,
            'odt': self._extract_docx_content,
            'odp': self._extract_pptx_content,
            'ods': self._extract_spreadsheet_content,
            'jpg': self._extract_image_content,
            'jpeg': self._extract_image_content,
            'png': self._extract_image_content,
            'gif': self._extract_image_content,
            'bmp': self._extract_image_content,
            'tiff': self._extract_image_content,
            'webp': self._extract_image_content,
            'svg': self._extract_image_content,
            'py': self._extract_text_content,
            'js': self._extract_text_content,
            'ts': self._extract_text_content,
            'java': self._extract_text_content,
            'cpp': self._extract_text_content,
            'c': self._extract_text_content,
            'go': self._extract_text_content,
            'rs': self._extract_text_content,
            'rb': self._extract_text_content,
            'php': self._extract_text_content
        }

        # Build supported formats based on global settings
        self.supported_formats: Dict[str, Any] = {}
        self._build_supported_formats()

        # Define modality mapping
        self.image_formats = {'.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff', '.webp', '.svg'}
        self.structured_formats = {'.csv', '.xlsx', '.xls', '.ods', '.json', '.xml'}

    async def process_file(self, file_path: str, file_extension: Optional[str] = None) -> Tuple[str, str, Dict[str, Any]]:
        """
        Extract text content from file and detect modality

        Args:
            file_path: Path to file
            file_extension: File extension (auto-detected if None)

        Returns:
            (content, modality, metadata) tuple where:
            - content: Extracted text content
            - modality: "text", "structured", or "image"
            - metadata: Extraction metadata (row counts, truncation info, etc.)
        """
        self._ensure_supported_formats()

        if file_extension is None:
            file_extension = Path(file_path).suffix.lower()
        else:
            file_extension = file_extension.lower()

        if file_extension not in self.supported_formats:
            return (
                f"Unsupported file format: {file_extension}",
                "text",
                {
                    "error": "unsupported_format",
                    "file_extension": file_extension
                }
            )

        try:
            processor = self.supported_formats[file_extension]
            content, metadata = processor(file_path)
            metadata = metadata or {}
            metadata.setdefault("file_extension", file_extension)

            modality = self._detect_modality(file_extension)
            metadata.setdefault("modality_detected", modality)

            return content, modality, metadata

        except Exception as exc:
            return (
                f"Error processing {file_extension} file: {str(exc)}",
                "text",
                {
                    "error": str(exc),
                    "file_extension": file_extension,
                    "summary_strategy": "extraction_error"
                }
            )

    def _detect_modality(self, file_extension: str) -> str:
        """Detect document modality based on file extension"""
        if file_extension in self.image_formats:
            return "image"
        if file_extension in self.structured_formats:
            return "structured"
        return "text"

    def _build_supported_formats(self) -> None:
        """Build handler mapping from the latest settings."""
        configured_formats = set()

        supported_file_formats = getattr(self.settings, "supported_file_formats", []) or []
        configured_formats.update(fmt.lower().lstrip('.') for fmt in supported_file_formats)

        extraction_settings = getattr(self.settings, "document_extraction", None)
        if extraction_settings and getattr(extraction_settings, "supported_formats", None):
            configured_formats.update(fmt.lower().lstrip('.') for fmt in extraction_settings.supported_formats)

        self.supported_formats = {}
        for fmt in configured_formats:
            handler = self._format_mapping.get(fmt)
            if handler:
                self.supported_formats[f'.{fmt}'] = handler

    def _ensure_supported_formats(self) -> None:
        """Refresh supported formats if settings have changed."""
        from src.core.config.settings import get_settings
        current_settings = get_settings()
        if id(current_settings) != self._settings_signature:
            self.settings = current_settings
            self._settings_signature = id(current_settings)
            self._build_supported_formats()

    # --------------------------------------------------------------------- #
    # Extraction helpers
    # --------------------------------------------------------------------- #

    def _extract_pdf_content(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text from PDF using PyMuPDF (fallback to PyPDF)"""
        try:
            content: List[str] = []
            metadata: Dict[str, Any] = {}

            doc = fitz.open(file_path)
            metadata["page_count"] = len(doc)

            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text()

                if text.strip():
                    content.append(f"--- Page {page_num + 1} ---\n{text}")
                else:
                    try:
                        pix = page.get_pixmap()
                        img_data = pix.tobytes("ppm")
                        with tempfile.NamedTemporaryFile(suffix=".ppm", delete=False) as temp_img:
                            temp_img.write(img_data)
                            temp_img.flush()

                            content.append(
                                f"--- Page {page_num + 1} (Image-based) ---\n"
                                "[Image content detected - AI analysis will be performed separately]"
                            )
                            os.unlink(temp_img.name)
                    except Exception:
                        content.append(f"--- Page {page_num + 1} ---\n[Unable to extract content from this page]")

            doc.close()
            metadata["summary_strategy"] = "pdf_native"
            result = '\n\n'.join(content) if content else "No readable content found in PDF"
            return result, metadata

        except Exception:
            try:
                content = []
                with open(file_path, 'rb') as file:
                    reader = pypdf.PdfReader(file)
                    metadata = {
                        "page_count": len(reader.pages),
                        "summary_strategy": "pdf_pypdf_fallback"
                    }
                    for page_num, page in enumerate(reader.pages):
                        text = page.extract_text()
                        if text and text.strip():
                            content.append(f"--- Page {page_num + 1} ---\n{text}")

                result = '\n\n'.join(content) if content else "No readable content found in PDF"
                return result, metadata
            except Exception as exc:
                return (
                    f"Error extracting PDF content: {str(exc)}",
                    {
                        "error": str(exc),
                        "summary_strategy": "pdf_error"
                    }
                )

    def _extract_docx_content(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text from Word documents"""
        try:
            doc = Document(file_path)
            content: List[str] = []
            paragraph_count = 0
            table_count = 0

            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    content.append(paragraph.text)
                    paragraph_count += 1

            for table in doc.tables:
                table_count += 1
                table_content: List[str] = []
                for row in table.rows:
                    row_content = [cell.text.strip() for cell in row.cells]
                    table_content.append(' | '.join(row_content))

                if table_content:
                    content.append('\n--- Table ---\n' + '\n'.join(table_content))

            text = '\n\n'.join(content) if content else "No readable content found in document"
            metadata = {
                "paragraph_count": paragraph_count,
                "table_count": table_count,
                "word_count": len(text.split()) if text else 0,
                "summary_strategy": "docx_native"
            }

            return text, metadata
        except Exception as exc:
            return (
                f"Error extracting DOCX content: {str(exc)}",
                {
                    "error": str(exc),
                    "summary_strategy": "docx_error"
                }
            )

    def _extract_pptx_content(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text from PowerPoint presentations"""
        try:
            prs = Presentation(file_path)
            content: List[str] = []
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_content = [f"--- Slide {slide_num} ---"]
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_content.append(shape.text)

                if len(slide_content) > 1:
                    content.append('\n'.join(slide_content))

            metadata = {
                "slide_count": len(prs.slides),
                "summary_strategy": "pptx_native"
            }
            result = '\n\n'.join(content) if content else "No readable content found in presentation"
            return result, metadata
        except Exception as exc:
            return (
                f"Error extracting PPTX content: {str(exc)}",
                {
                    "error": str(exc),
                    "summary_strategy": "pptx_error"
                }
            )

    def _extract_csv_content(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract and summarize CSV data with sampling"""
        try:
            total_rows = self._estimate_csv_row_count(file_path)

            sample_df = pd.read_csv(
                file_path,
                nrows=MAX_TABULAR_SAMPLE_ROWS,
                low_memory=False
            )

            row_count = total_rows if total_rows is not None else len(sample_df)
            truncated = row_count > len(sample_df)

            return self._summarize_tabular_dataframe(
                df=sample_df,
                row_count=row_count,
                truncated=truncated,
                source_label="CSV Data Summary",
                strategy="csv_head_sample"
            )
        except Exception as exc:
            return (
                f"Error extracting CSV content: {str(exc)}",
                {
                    "error": str(exc),
                    "summary_strategy": "csv_error"
                }
            )

    def _extract_spreadsheet_content(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract and summarize spreadsheet data (Excel/ODS)"""
        try:
            total_rows = self._estimate_spreadsheet_rows(file_path)
            sample_df = pd.read_excel(
                file_path,
                engine=None,
                nrows=MAX_TABULAR_SAMPLE_ROWS
            )

            row_count = total_rows if total_rows is not None else len(sample_df)
            truncated = row_count > len(sample_df)

            return self._summarize_tabular_dataframe(
                df=sample_df,
                row_count=row_count,
                truncated=truncated,
                source_label="Spreadsheet Data Summary",
                strategy="spreadsheet_head_sample"
            )
        except Exception as exc:
            return (
                f"Error extracting spreadsheet content: {str(exc)}",
                {
                    "error": str(exc),
                    "summary_strategy": "spreadsheet_error"
                }
            )

    def _extract_text_content(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract content from plain text-like files"""
        try:
            encodings = ['utf-8', 'latin-1', 'cp1252']
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as file:
                        text = file.read()
                        return text, {
                            "character_count": len(text),
                            "summary_strategy": "text_native"
                        }
                except UnicodeDecodeError:
                    continue

            return (
                "Error: Could not read text file with standard encodings",
                {
                    "error": "decode_error",
                    "summary_strategy": "text_error"
                }
            )
        except Exception as exc:
            return (
                f"Error extracting text content: {str(exc)}",
                {
                    "error": str(exc),
                    "summary_strategy": "text_error"
                }
            )

    def _extract_image_content(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract metadata from images (content handled by vision model later)"""
        try:
            doc = fitz.open(file_path)
            page = doc.load_page(0)
            pix = page.get_pixmap()

            width, height = pix.width, pix.height
            metadata = {
                "width": width,
                "height": height,
                "summary_strategy": "image_stub"
            }

            content = [
                "Image Analysis:",
                f"Dimensions: {width} x {height} pixels",
                f"Format: {Path(file_path).suffix.upper()}",
                "\n[Image detected - AI analysis will be performed separately]"
            ]

            # Attempt simple heuristics via PIL for additional metadata
            try:
                img_data = pix.tobytes("ppm")
                with tempfile.NamedTemporaryFile(suffix=".ppm", delete=False) as temp_img:
                    temp_img.write(img_data)
                    temp_img.flush()
                    Image.open(temp_img.name)  # Ensures image readable
                    os.unlink(temp_img.name)
            except Exception:
                metadata["analysis_warning"] = "basic_image_analysis_failed"

            doc.close()
            return '\n'.join(content), metadata
        except Exception as exc:
            return (
                f"Error analyzing image: {str(exc)}",
                {
                    "error": str(exc),
                    "summary_strategy": "image_error"
                }
            )

    # --------------------------------------------------------------------- #
    # Tabular helpers
    # --------------------------------------------------------------------- #

    def _summarize_tabular_dataframe(
        self,
        df: pd.DataFrame,
        row_count: int,
        truncated: bool,
        source_label: str,
        strategy: str
    ) -> Tuple[str, Dict[str, Any]]:
        preview_rows = min(TABULAR_PREVIEW_ROWS, len(df))

        metadata: Dict[str, Any] = {
            "row_count": int(row_count),
            "sampled_rows": int(len(df)),
            "preview_rows": int(preview_rows),
            "column_count": int(len(df.columns)),
            "columns": [str(col) for col in df.columns[:MAX_COLUMN_NAMES]],
            "truncated": bool(truncated),
            "summary_strategy": strategy
        }

        if len(df.columns) > MAX_COLUMN_NAMES:
            metadata["columns_truncated"] = True

        content: List[str] = [
            source_label,
            f"Total rows (approx): {row_count}",
            f"Columns: {len(df.columns)}"
        ]

        if len(df.columns) > 0:
            columns_preview = ', '.join(str(col) for col in df.columns[:MAX_COLUMN_NAMES])
            if len(df.columns) > MAX_COLUMN_NAMES:
                columns_preview += ", …"
            content.append(f"Column names: {columns_preview}")

        if truncated:
            content.append(
                f"\n⚠️ Large dataset detected. Only the first {len(df)} rows were processed for embeddings."
            )

        if preview_rows > 0:
            content.append(f"\n--- First {preview_rows} rows ---")
            content.append(df.head(preview_rows).to_string(index=False))

        numeric_cols = df.select_dtypes(include=['number']).columns.tolist()
        if numeric_cols:
            cols_for_stats = [str(col) for col in numeric_cols[:NUMERIC_SUMMARY_MAX_COLS]]
            metadata["numeric_columns_profiled"] = cols_for_stats
            if len(numeric_cols) > NUMERIC_SUMMARY_MAX_COLS:
                metadata["numeric_columns_truncated"] = True
            content.append("\n--- Numeric Summary ---")
            content.append(df[numeric_cols[:NUMERIC_SUMMARY_MAX_COLS]].describe().to_string())

        metadata["notes"] = (
            f"Processed first {len(df)} rows out of approximately {row_count} for embeddings."
            if truncated else
            f"Processed all {row_count} rows for embeddings."
        )

        return '\n'.join(content), metadata

    def _estimate_spreadsheet_rows(self, file_path: str) -> Optional[int]:
        try:
            workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
            sheet = workbook.active
            count = max((sheet.max_row or 1) - 1, 0)
            workbook.close()
            return count
        except Exception:
            return None

    def _estimate_csv_row_count(self, file_path: str) -> Optional[int]:
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                # subtract header row if present
                return max(sum(1 for _ in f) - 1, 0)
        except Exception:
            return None


# Global instance
document_extractor = DocumentExtractor()
